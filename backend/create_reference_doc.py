from docx import Document

document = Document()

document.add_paragraph(
    "Supervised learning uses labeled examples to learn relationships "
    "between inputs and known outputs."
)

document.add_paragraph(
    "Unsupervised learning discovers hidden patterns in data "
    "without predefined labels."
)

document.save("reference_submission.docx")

print("Created reference_submission.docx")
