from docx import Document

document = Document()

document.add_paragraph(
    "Machine learning is a branch of artificial intelligence that allows computers to learn patterns from data."
)

document.add_paragraph(
    "Supervised learning uses labeled examples to learn relationships between inputs and known outputs."
)

document.add_paragraph(
    "Unsupervised learning discovers hidden patterns in data without predefined labels."
)

document.save("clean_reference.docx")

print("Created clean_reference.docx")
