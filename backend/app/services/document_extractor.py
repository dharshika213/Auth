import os
import re

import pymupdf
from docx import Document
from pptx import Presentation


def clean_extracted_text(text: str) -> str:
    """
    Normalize whitespace for PDF/PPTX extracted text.

    This function intentionally does NOT attempt to guess
    missing word boundaries.
    """

    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")
    text = text.replace("\t", " ")

    # Collapse repeated spaces.
    text = re.sub(r"[ \t]+", " ", text)

    # Remove excessive blank lines.
    text = re.sub(r"\n\s*\n+", "\n\n", text)

    return text.strip()


def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []

    with pymupdf.open(file_path) as document:
        for page in document:
            page_text = page.get_text(
                "text",
                sort=True,
            )

            if page_text:
                text_parts.append(page_text)

    return clean_extracted_text(
        "\n".join(text_parts)
    )


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text directly from DOCX.

    IMPORTANT:
    python-docx already preserves the actual text and word
    spacing from the document. Therefore we do NOT run the
    generic cleanup function on DOCX text.
    """

    document = Document(file_path)

    text_parts = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text

        if paragraph_text.strip():
            text_parts.append(
                paragraph_text
            )

    # Extract text from tables.
    for table in document.tables:
        for row in table.rows:

            row_text = []

            for cell in row.cells:
                cell_text = cell.text

                if cell_text.strip():
                    row_text.append(
                        cell_text.strip()
                    )

            if row_text:
                text_parts.append(
                    " | ".join(row_text)
                )

    # Do NOT call clean_extracted_text() here.
    return "\n".join(text_parts).strip()


def extract_text_from_pptx(file_path: str) -> str:
    presentation = Presentation(file_path)

    text_parts = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            shape_text = shape.text.strip()

            if shape_text:
                text_parts.append(
                    shape_text
                )

    return clean_extracted_text(
        "\n".join(text_parts)
    )


def extract_text(file_path: str) -> str:
    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".pptx":
        return extract_text_from_pptx(file_path)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )

